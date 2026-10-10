"""
importar_entrada.py

Lê a planilha de entrada de mercadoria (a que "Gerar planilha de entrada de
mercadoria" cria, já preenchida com o que chegou) e aplica no banco:
atualiza custo e preço de venda, dá entrada no estoque, cadastra produto
novo quando for o caso, e sincroniza de volta a margem por categoria e a
taxa de cartão se tiverem sido editadas na aba "Margens".

Sempre em duas etapas — nunca grava sem passar pela primeira:

    pre_visualizar(caminho)   -> só lê e monta o que SERIA feito
    aplicar(caminho)          -> faz um backup do banco e grava de verdade

Uso pela interface: Arquivo -> Importar entrada de mercadoria.
"""
from __future__ import annotations

import os
import sqlite3
from datetime import datetime

import database as db
import operacoes as op
from precificacao import carregar_margens, salvar_margens, config_precificacao

try:
    import openpyxl
except ImportError:
    openpyxl = None

COLUNAS_LIDAS = 17  # mesmo tanto de colunas de precificacao.COLUNAS_ENTRADA
LINHA_CABECALHO = 4


def _checar_openpyxl():
    if openpyxl is None:
        raise RuntimeError('Este recurso precisa do pacote openpyxl. Rode: pip install openpyxl')


def _numero(valor):
    return float(valor) if isinstance(valor, (int, float)) else None


def _ler_margens_da_planilha(wb):
    """Lê a aba Margens, se existir, pra sincronizar de volta no config.json.
    Devolve (margens_por_categoria, taxa_cartao) ou (None, None) se a aba
    não existir ou não tiver nada preenchido."""
    if 'Margens' not in wb.sheetnames:
        return None, None
    ws = wb['Margens']
    margens = {}
    linha = 5
    vazias_seguidas = 0
    while vazias_seguidas < 2:
        cat = ws.cell(row=linha, column=1).value
        if not cat or 'taxa' in str(cat).lower() or 'arredondamento' in str(cat).lower():
            vazias_seguidas += 1
            linha += 1
            continue
        vazias_seguidas = 0
        valor = _numero(ws.cell(row=linha, column=2).value)
        if valor is not None:
            margens[str(cat).strip()] = valor
        linha += 1
        if linha > 200:  # trava de segurança
            break

    taxa = None
    for r in range(1, ws.max_row + 1):
        rotulo = ws.cell(row=r, column=1).value
        if rotulo and 'taxa média da maquininha' in str(rotulo).lower():
            taxa = _numero(ws.cell(row=r, column=2).value)
            break
    return (margens or None), taxa


def _ler_linhas_entrada(wb):
    if 'Entrada' not in wb.sheetnames:
        raise ValueError('Essa planilha não tem uma aba "Entrada" — não parece ser o modelo gerado pelo sistema.')
    ws = wb['Entrada']
    linhas = []
    linha = LINHA_CABECALHO + 1
    vazias_seguidas = 0
    while vazias_seguidas < 3 and linha < ws.max_row + 3:
        valores = [ws.cell(row=linha, column=c).value for c in range(1, COLUNAS_LIDAS + 1)]
        nome = valores[1]
        if not nome or not str(nome).strip():
            vazias_seguidas += 1
            linha += 1
            continue
        vazias_seguidas = 0
        codigo, _nome, categoria, unidade, qtd, _custo_total, custo_unit, _margem, _preco_cheio, \
            _preco_sugerido, preco_atual, preco_final = valores[:12]
        linhas.append({
            'linha_planilha': linha,
            'codigo_barras': str(codigo).strip() if codigo not in (None, '') else None,
            'nome': str(nome).strip(),
            'categoria': str(categoria).strip() if categoria not in (None, '') else None,
            'unidade': str(unidade).strip().upper() if unidade in ('UN', 'KG') else 'UN',
            'quantidade': _numero(qtd) or 0.0,
            'custo_unitario': _numero(custo_unit) or 0.0,
            'preco_atual_na_planilha': _numero(preco_atual),
            'preco_final': _numero(preco_final),
        })
        linha += 1
    return linhas


def _buscar_produto_existente(item):
    if item['codigo_barras']:
        p = op.buscar_por_codigo(item['codigo_barras'])
        if p:
            return p
    conn = db.get_conn()
    try:
        row = conn.execute(
            'SELECT * FROM produtos WHERE UPPER(TRIM(nome)) = ? AND ativo = 1',
            (item['nome'].strip().upper(),)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def _montar_plano(caminho_planilha):
    """Lê a planilha e monta a lista de ações, sem tocar no banco.
    Devolve {'itens': [...], 'margens_novas': ..., 'taxa_nova': ...,
    'total_custo': ..., 'total_itens': ...}"""
    _checar_openpyxl()
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    linhas = _ler_linhas_entrada(wb)
    margens_novas, taxa_nova = _ler_margens_da_planilha(wb)

    itens = []
    total_custo = 0.0
    for item in linhas:
        existente = _buscar_produto_existente(item)
        preco_final = item['preco_final'] if item['preco_final'] and item['preco_final'] > 0 else None
        avisos = []
        if item['custo_unitario'] <= 0:
            avisos.append('sem custo informado')
        if item['quantidade'] <= 0:
            avisos.append('quantidade zerada — não vai dar entrada em estoque')
        if preco_final and item['custo_unitario'] and preco_final <= item['custo_unitario']:
            avisos.append('PREJUÍZO: preço final menor ou igual ao custo')

        if existente:
            acao = 'atualizar'
            plano = {
                'produto_id': existente['id'],
                'nome': existente['nome'],
                'preco_custo_antes': existente['preco_custo'],
                'preco_venda_antes': existente['preco_venda'],
                'estoque_antes': existente['quantidade_estoque'],
                'estoque_depois': round((existente['quantidade_estoque'] or 0) + item['quantidade'], 3),
                'tinha_codigo': bool(existente['codigo_barras'] and str(existente['codigo_barras']).strip()),
            }
        else:
            acao = 'cadastrar'
            if not preco_final:
                avisos.append('produto novo sem preço final definido — não será cadastrado')
            plano = {'produto_id': None, 'nome': item['nome']}

        total_custo += item['custo_unitario'] * item['quantidade']
        itens.append({
            **item, 'acao': acao, 'plano': plano, 'avisos': avisos,
            'preco_custo_novo': item['custo_unitario'],
            'preco_venda_novo': preco_final,
        })

    return {
        'itens': itens,
        'margens_novas': margens_novas,
        'taxa_nova': taxa_nova,
        'total_custo': round(total_custo, 2),
        'total_itens': len(itens),
    }


def pre_visualizar(caminho_planilha):
    """Monta o plano de importação pra mostrar antes de gravar qualquer coisa."""
    return _montar_plano(caminho_planilha)


def _backup_antes_de_gravar():
    """Backup consistente (API do SQLite) antes de qualquer gravação."""
    pasta = os.path.join(os.path.dirname(os.path.abspath(db.DB_PATH)), 'backups')
    os.makedirs(pasta, exist_ok=True)
    destino = os.path.join(pasta, f'mercadinho_antes_entrada_{datetime.now().strftime("%Y%m%d_%H%M%S")}.db')
    origem = sqlite3.connect(db.DB_PATH)
    copia = sqlite3.connect(destino)
    try:
        origem.backup(copia)
        if copia.execute('PRAGMA integrity_check').fetchone()[0] != 'ok':
            raise RuntimeError('Backup não íntegro antes da entrada de mercadoria.')
    finally:
        origem.close()
        copia.close()
    return destino


def aplicar(caminho_planilha, plano=None):
    """Grava de verdade: atualiza/cadastra produtos, dá entrada em estoque
    e sincroniza margens. Faz backup do banco antes de qualquer gravação.

    Devolve um resumo com o que foi feito, pra interface mostrar."""
    import gestao as _gestao
    _gestao.permitir('estoque')
    plano = plano or _montar_plano(caminho_planilha)
    backup = _backup_antes_de_gravar()

    atualizados, cadastrados, ignorados = [], [], []
    for item in plano['itens']:
        if item['quantidade'] <= 0 and item['acao'] == 'atualizar' and item['custo_unitario'] <= 0:
            ignorados.append(item['nome'])
            continue
        try:
            if item['acao'] == 'atualizar':
                pid = item['plano']['produto_id']
                kwargs = {}
                if item['custo_unitario'] > 0:
                    kwargs['preco_custo'] = item['custo_unitario']
                if item['preco_venda_novo']:
                    kwargs['preco_venda'] = item['preco_venda_novo']
                if item['categoria']:
                    kwargs['categoria'] = item['categoria']
                if item['codigo_barras'] and not item['plano'].get('tinha_codigo'):
                    kwargs['codigo_barras'] = item['codigo_barras']
                if kwargs:
                    op.atualizar_produto(pid, **kwargs)
                if item['quantidade'] > 0:
                    op.ajustar_estoque(pid, item['quantidade'], 'entrada',
                                       f'Entrada de mercadoria {datetime.now().strftime("%d/%m/%Y")}')
                atualizados.append(item['nome'])
            else:
                if not item['preco_venda_novo']:
                    ignorados.append(item['nome'])
                    continue
                op.criar_produto(
                    item['nome'], codigo_barras=item['codigo_barras'], categoria=item['categoria'],
                    unidade=item['unidade'], preco_custo=item['custo_unitario'],
                    preco_venda=item['preco_venda_novo'], quantidade_estoque=item['quantidade'],
                    estoque_minimo=2 if item['unidade'] == 'KG' else 5)
                cadastrados.append(item['nome'])
        except op.ErroOperacao as exc:
            ignorados.append(f"{item['nome']} ({exc})")

    margens_sincronizadas = False
    if plano['margens_novas'] or plano['taxa_nova'] is not None:
        atuais = carregar_margens()
        atuais.update(plano['margens_novas'] or {})
        salvar_margens(atuais)
        if plano['taxa_nova'] is not None:
            from impressora import salvar_config
            salvar_config({'taxa_cartao_referencia': plano['taxa_nova']})
        margens_sincronizadas = True

    return {
        'backup': backup,
        'atualizados': atualizados,
        'cadastrados': cadastrados,
        'ignorados': ignorados,
        'margens_sincronizadas': margens_sincronizadas,
        'total_custo': plano['total_custo'],
    }


if __name__ == '__main__':
    import sys
    if len(sys.argv) < 2:
        print('Uso: python importar_entrada.py caminho/da/planilha_preenchida.xlsx')
        sys.exit(1)
    db.init_db()
    plano = pre_visualizar(sys.argv[1])
    print(f"{plano['total_itens']} linha(s) na planilha, custo total R$ {plano['total_custo']:.2f}\n")
    for item in plano['itens']:
        marca = ' | '.join(item['avisos']) if item['avisos'] else 'ok'
        print(f"  [{item['acao']:10}] {item['nome']:40} custo {item['custo_unitario']:.2f} -> venda {item['preco_venda_novo']}  ({marca})")
    if plano['margens_novas']:
        print('\nMargens que serão sincronizadas:', plano['margens_novas'])
    resposta = input('\nAplicar de verdade? (s/N) ').strip().lower()
    if resposta == 's':
        resultado = aplicar(sys.argv[1], plano)
        print(f"Atualizados: {len(resultado['atualizados'])} | Cadastrados: {len(resultado['cadastrados'])} | "
              f"Ignorados: {len(resultado['ignorados'])}")
        print(f"Backup do banco em: {resultado['backup']}")
    else:
        print('Nada foi gravado.')
