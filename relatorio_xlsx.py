"""
relatorio_xlsx.py

Exporta todas as análises de vendas pra uma planilha .xlsx formatada — no
mesmo espírito do catálogo de produtos: cabeçalhos coloridos, valores em
R$ de verdade (não texto), filtros, painéis congelados e gráficos nativos
do Excel, que continuam funcionando e podem ser reposicionados na mão.

Abas geradas:
    Resumo                  — cartões do período + alertas de cadastro
    Evolução por Dia        — série diária com gráfico de linha
    Evolução por Semana     — série semanal com gráfico de barras
    Evolução por Mês        — série mensal com gráfico de barras
    Dia da Semana           — total e média por dia da semana, com gráfico
    O que vende em cada dia — top produtos de cada dia da semana
    Horário                 — movimento por hora, com gráfico
    Top Produtos            — ranking por quantidade e receita
    Curva ABC               — classificação A/B/C dos produtos
    Categorias              — participação de cada categoria
    Formas de Pagamento     — quanto entrou em cada forma e taxa paga
    Tipo de Dia              — feriado, véspera, dia de pagamento, fim de semana
    Clientes                — ranking de clientes cadastrados

Uso pela interface: Análises -> Exportar relatório (.xlsx).
Uso direto: python relatorio_xlsx.py [saida.xlsx] [inicio AAAA-MM-DD] [fim AAAA-MM-DD]
"""
from __future__ import annotations

import sys
from datetime import datetime

import analises as an
import database as db

try:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

VERDE_ESCURO = '1B5240'
VERDE = '2E7D32'
VERDE_CLARO = 'E4EFE7'
AMARELO = 'F3B22B'
AMARELO_CLARO = 'FDF2DC'
VERMELHO = 'C23B2E'
VERMELHO_CLARO = 'FBE8E4'
CINZA = '6B695F'
CABECALHO = 'C8E6C9'

MOEDA = 'R$ #,##0.00'
NUMERO = '#,##0.000'
INTEIRO = '#,##0'
PCT = '0.0"%"'


# ----------------------------------------------------------------------
# Auxiliares de formatação
# ----------------------------------------------------------------------
def _titulo_aba(ws, titulo, subtitulo=''):
    ws['A1'] = titulo
    ws['A1'].font = Font(size=15, bold=True, color=VERDE_ESCURO)
    if subtitulo:
        ws['A2'] = subtitulo
        ws['A2'].font = Font(size=9, italic=True, color=CINZA)
    return 4 if subtitulo else 3


def _cabecalho(ws, linha, titulos, larguras=None, cor=CABECALHO):
    for i, titulo in enumerate(titulos, start=1):
        cel = ws.cell(row=linha, column=i, value=titulo)
        cel.font = Font(bold=True, color=VERDE_ESCURO)
        cel.fill = PatternFill('solid', fgColor=cor)
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if larguras:
        for i, largura in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha].height = 26
    return linha + 1


def _formatar_coluna(ws, coluna, primeira, ultima, formato):
    letra = get_column_letter(coluna)
    for r in range(primeira, ultima + 1):
        ws[f'{letra}{r}'].number_format = formato


def _zebra(ws, primeira, ultima, n_colunas):
    for r in range(primeira, ultima + 1):
        if (r - primeira) % 2 == 1:
            for c in range(1, n_colunas + 1):
                ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor='FAFBF7')


def _grafico(ws, tipo, titulo, dados_col, cat_col, primeira, ultima, ancora, largura=24, altura=9):
    """Adiciona um gráfico nativo do Excel apontando pras células da aba."""
    grafico = LineChart() if tipo == 'linha' else BarChart()
    if tipo != 'linha':
        grafico.type = 'col'
    grafico.title = titulo
    grafico.height = altura
    grafico.width = largura
    grafico.legend = None
    dados = Reference(ws, min_col=dados_col, min_row=primeira - 1, max_row=ultima)
    categorias = Reference(ws, min_col=cat_col, min_row=primeira, max_row=ultima)
    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)
    ws.add_chart(grafico, ancora)


# ----------------------------------------------------------------------
# Abas
# ----------------------------------------------------------------------
def _aba_resumo(wb, inicio, fim, rotulo):
    ws = wb.active
    ws.title = 'Resumo'
    ws['A1'] = 'FMG Mercadinho — Relatório de Vendas'
    ws['A1'].font = Font(size=18, bold=True, color=VERDE_ESCURO)
    ws['A2'] = f'Período: {rotulo}'
    ws['A2'].font = Font(size=11, bold=True, color=VERDE)
    ws['A3'] = f'Gerado em {datetime.now().strftime("%d/%m/%Y às %H:%M")}'
    ws['A3'].font = Font(size=9, italic=True, color=CINZA)

    r = an.resumo_periodo(inicio, fim)
    linha = 5
    cartoes = [
        ('Faturamento', r['faturamento'], MOEDA),
        ('Número de vendas', r['numero_vendas'], INTEIRO),
        ('Ticket médio', r['ticket_medio'], MOEDA),
        ('Itens vendidos', r['itens_vendidos'], NUMERO),
        ('Lucro estimado', r['lucro_estimado'], MOEDA),
        ('Descontos dados', r['descontos'], MOEDA),
        ('Taxas de maquininha/voucher', r['taxas'], MOEDA),
        ('Dias com movimento', r['dias_com_venda'], INTEIRO),
        ('Média por dia aberto', r['media_por_dia'], MOEDA),
        ('Vendas por dia aberto', r['vendas_por_dia'], '#,##0.0'),
    ]
    linha = _cabecalho(ws, linha, ['Indicador', 'Valor'], [38, 20])
    primeira = linha
    for nome, valor, formato in cartoes:
        ws.cell(row=linha, column=1, value=nome)
        cel = ws.cell(row=linha, column=2, value=valor)
        cel.number_format = formato
        cel.font = Font(bold=True, color=VERDE_ESCURO)
        linha += 1
    _zebra(ws, primeira, linha - 1, 2)

    # Melhor e pior dia da semana
    melhor, pior = an.melhor_pior_dia(inicio, fim)
    linha += 1
    if melhor:
        ws.cell(row=linha, column=1, value='Melhor dia da semana').font = Font(bold=True)
        ws.cell(row=linha, column=2, value=f"{melhor['nome']} — {melhor['numero_vendas']} venda(s)")
        linha += 1
        ws.cell(row=linha, column=1, value='Dia mais fraco').font = Font(bold=True)
        ws.cell(row=linha, column=2, value=f"{pior['nome']} — {pior['numero_vendas']} venda(s)")
        linha += 2

    # Diagnóstico do banco usado
    d = an.diagnostico_banco()
    ws.cell(row=linha, column=1, value='BANCO DE DADOS ANALISADO').font = Font(bold=True, color=VERDE_ESCURO)
    linha += 1
    for nome, valor in [
        ('Produtos ativos', d['produtos_ativos']),
        ('Clientes cadastrados', d['clientes_ativos']),
        ('Total de vendas no banco', d['total_vendas']),
        ('Primeira venda registrada', d['primeira_venda'] or '—'),
        ('Última venda registrada', d['ultima_venda'] or '—'),
    ]:
        ws.cell(row=linha, column=1, value=nome)
        ws.cell(row=linha, column=2, value=valor)
        linha += 1
    if not d['tem_taxa_paga']:
        ws.cell(row=linha, column=1,
                value='Obs.: este banco não guarda a taxa de maquininha; as taxas aparecem como zero.')
        ws.cell(row=linha, column=1).font = Font(italic=True, color=CINZA)
        linha += 1

    # Alertas de cadastro
    q = an.qualidade_dados()
    linha += 1
    ws.cell(row=linha, column=1, value='PONTOS DE ATENÇÃO NO CADASTRO').font = Font(bold=True, color=VERMELHO)
    linha += 1
    alertas = []
    if q['sem_preco_custo']:
        alertas.append((f"{q['sem_preco_custo']} produto(s) sem preço de custo preenchido",
                        'Enquanto o custo estiver zerado, o "lucro estimado" sai praticamente igual ao faturamento.'))
    if q['sem_categoria']:
        alertas.append((f"{q['sem_categoria']} produto(s) sem categoria",
                        'Sem categoria eles caem no grupo "Sem categoria" nas análises.'))
    if q['categorias_fora_padrao']:
        extras = ', '.join(f"{c['categoria']} ({c['n']})" for c in q['categorias_fora_padrao'])
        alertas.append((f'Categorias fora do padrão: {extras}', 'Vale padronizar pra não dividir o mesmo grupo em dois.'))
    if q['sem_codigo_balanca']:
        alertas.append((f"{q['sem_codigo_balanca']} item(ns) de balança sem código pra bipar",
                        'Dá pra resolver em Arquivo → Gerar catálogo atualizado.'))
    if q['sem_codigo_un']:
        alertas.append((f"{q['sem_codigo_un']} produto(s) por unidade sem código de barras", ''))
    if not alertas:
        ws.cell(row=linha, column=1, value='Nenhum ponto de atenção — cadastro em ordem.')
    for texto, detalhe in alertas:
        ws.cell(row=linha, column=1, value=texto).font = Font(bold=True, color=VERMELHO)
        ws.cell(row=linha, column=1).fill = PatternFill('solid', fgColor=VERMELHO_CLARO)
        if detalhe:
            ws.cell(row=linha, column=2, value=detalhe).font = Font(italic=True, color=CINZA)
        linha += 1

    ws.column_dimensions['A'].width = 52
    ws.column_dimensions['B'].width = 64
    return ws


def _aba_serie(wb, nome_aba, granularidade, titulo, inicio, fim, rotulo, tipo_grafico):
    dados = an.serie_temporal(granularidade, inicio, fim)
    ws = wb.create_sheet(nome_aba)
    linha = _titulo_aba(ws, titulo, f'Período: {rotulo}')
    linha = _cabecalho(ws, linha, ['Período', 'Vendas', 'Faturamento', 'Ticket médio', 'Lucro estimado'],
                       [22, 12, 16, 15, 16])
    primeira = linha
    for item in dados:
        ws.cell(row=linha, column=1, value=item['rotulo'])
        ws.cell(row=linha, column=2, value=item['numero_vendas'])
        ws.cell(row=linha, column=3, value=item['faturamento'])
        ws.cell(row=linha, column=4, value=item['ticket_medio'])
        ws.cell(row=linha, column=5, value=item['lucro_estimado'])
        linha += 1
    ultima = linha - 1
    if dados:
        for col in (3, 4, 5):
            _formatar_coluna(ws, col, primeira, ultima, MOEDA)
        _formatar_coluna(ws, 2, primeira, ultima, INTEIRO)
        _zebra(ws, primeira, ultima, 5)
        ws.freeze_panes = ws.cell(row=primeira, column=1)
        _grafico(ws, tipo_grafico, f'Faturamento — {titulo}', 3, 1, primeira, ultima, f'G{primeira}')
    else:
        ws.cell(row=primeira, column=1, value='Nenhuma venda no período.')
    return ws


def _aba_dia_semana(wb, inicio, fim, rotulo):
    dados = an.vendas_por_dia_semana(inicio, fim)
    ws = wb.create_sheet('Dia da Semana')
    linha = _titulo_aba(ws, 'Vendas por dia da semana', f'Período: {rotulo}')
    linha = _cabecalho(ws, linha,
                       ['Dia', 'Vendas', 'Faturamento', 'Ticket médio', 'Dias abertos', 'Média por dia'],
                       [14, 12, 16, 15, 14, 16])
    primeira = linha
    for item in dados:
        ws.cell(row=linha, column=1, value=item['nome'])
        ws.cell(row=linha, column=2, value=item['numero_vendas'])
        ws.cell(row=linha, column=3, value=item['faturamento'])
        ws.cell(row=linha, column=4, value=item['ticket_medio'])
        ws.cell(row=linha, column=5, value=item['ocorrencias'])
        ws.cell(row=linha, column=6, value=item['media_por_ocorrencia'])
        linha += 1
    ultima = linha - 1
    for col in (3, 4, 6):
        _formatar_coluna(ws, col, primeira, ultima, MOEDA)
    _zebra(ws, primeira, ultima, 6)

    melhor = max(dados, key=lambda d: d['faturamento'])
    for r in range(primeira, ultima + 1):
        if ws.cell(row=r, column=1).value == melhor['nome'] and melhor['faturamento'] > 0:
            for c in range(1, 7):
                ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=AMARELO_CLARO)
                ws.cell(row=r, column=c).font = Font(bold=True, color=VERDE_ESCURO)

    _grafico(ws, 'barra', 'Faturamento por dia da semana', 3, 1, primeira, ultima, f'H{primeira}')
    linha = ultima + 2
    ws.cell(row=linha, column=1,
            value='"Média por dia" divide o faturamento pelo número de vezes que esse dia da semana teve movimento — '
                  'é a comparação mais justa quando o período tem, por exemplo, 2 domingos e 1 sábado.')
    ws.cell(row=linha, column=1).font = Font(italic=True, color=CINZA)
    return ws


def _aba_produtos_por_dia(wb, inicio, fim, rotulo):
    ws = wb.create_sheet('O que vende em cada dia')
    linha = _titulo_aba(ws, 'O que mais vende em cada dia da semana', f'Período: {rotulo}')
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 42
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 24
    for dia in range(7):
        produtos = an.produtos_por_dia_semana(dia, inicio, fim, limite=10)
        cel = ws.cell(row=linha, column=1, value=an.NOMES_DIA_SEMANA[dia])
        cel.font = Font(bold=True, size=12, color='FFFFFF')
        cel.fill = PatternFill('solid', fgColor=VERDE)
        for c in range(2, 6):
            ws.cell(row=linha, column=c).fill = PatternFill('solid', fgColor=VERDE)
        linha += 1
        if not produtos:
            ws.cell(row=linha, column=2, value='Sem vendas nesse dia da semana no período').font = Font(italic=True, color=CINZA)
            linha += 2
            continue
        linha = _cabecalho(ws, linha, ['#', 'Produto', 'Quantidade', 'Receita', 'Categoria'])
        primeira = linha
        for i, p in enumerate(produtos, start=1):
            ws.cell(row=linha, column=1, value=i)
            ws.cell(row=linha, column=2, value=p['nome'])
            ws.cell(row=linha, column=3, value=round(p['qtd'], 3))
            ws.cell(row=linha, column=4, value=round(p['receita'], 2))
            ws.cell(row=linha, column=5, value=p['categoria'])
            linha += 1
        _formatar_coluna(ws, 3, primeira, linha - 1, NUMERO)
        _formatar_coluna(ws, 4, primeira, linha - 1, MOEDA)
        _zebra(ws, primeira, linha - 1, 5)
        linha += 1
    return ws


def _aba_horario(wb, inicio, fim, rotulo):
    dados = an.vendas_por_hora(inicio, fim)
    ws = wb.create_sheet('Horário')
    linha = _titulo_aba(ws, 'Movimento por horário', f'Período: {rotulo}')
    linha = _cabecalho(ws, linha, ['Hora', 'Vendas', 'Faturamento'], [12, 12, 16])
    primeira = linha
    for item in dados:
        ws.cell(row=linha, column=1, value=f"{item['hora']:02d}h")
        ws.cell(row=linha, column=2, value=item['numero_vendas'])
        ws.cell(row=linha, column=3, value=item['faturamento'])
        linha += 1
    ultima = linha - 1
    _formatar_coluna(ws, 3, primeira, ultima, MOEDA)
    _zebra(ws, primeira, ultima, 3)
    _grafico(ws, 'barra', 'Vendas por horário', 2, 1, primeira, ultima, f'E{primeira}')
    return ws


def _aba_top_produtos(wb, inicio, fim, rotulo):
    produtos = an.produtos_mais_vendidos(inicio, fim, limite=100)
    ws = wb.create_sheet('Top Produtos')
    linha = _titulo_aba(ws, 'Produtos mais vendidos', f'Período: {rotulo}')
    linha = _cabecalho(ws, linha, ['#', 'Produto', 'Categoria', 'Unid.', 'Quantidade', 'Receita', 'Nº de vendas'],
                       [6, 44, 22, 8, 14, 16, 14])
    primeira = linha
    for i, p in enumerate(produtos, start=1):
        ws.cell(row=linha, column=1, value=i)
        ws.cell(row=linha, column=2, value=p['nome'])
        ws.cell(row=linha, column=3, value=p['categoria'])
        ws.cell(row=linha, column=4, value=p['unidade'])
        ws.cell(row=linha, column=5, value=round(p['qtd'], 3))
        ws.cell(row=linha, column=6, value=round(p['receita'], 2))
        ws.cell(row=linha, column=7, value=p['n_vendas'])
        linha += 1
    ultima = linha - 1
    if produtos:
        _formatar_coluna(ws, 5, primeira, ultima, NUMERO)
        _formatar_coluna(ws, 6, primeira, ultima, MOEDA)
        _zebra(ws, primeira, ultima, 7)
        ws.auto_filter.ref = f'A{primeira - 1}:G{ultima}'
        ws.freeze_panes = ws.cell(row=primeira, column=1)
    else:
        ws.cell(row=primeira, column=2, value='Nenhuma venda no período.')
    return ws


def _aba_curva_abc(wb, inicio, fim, rotulo):
    produtos = an.curva_abc(inicio, fim)
    ws = wb.create_sheet('Curva ABC')
    linha = _titulo_aba(
        ws, 'Curva ABC de produtos',
        'Classe A = os produtos que somam ~80% da receita; B = até ~95%; C = a cauda longa. '
        f'Período: {rotulo}')
    linha = _cabecalho(ws, linha, ['Classe', 'Produto', 'Categoria', 'Quantidade', 'Receita', '% da receita', '% acumulado'],
                       [9, 44, 22, 14, 16, 14, 14])
    primeira = linha
    cores = {'A': VERDE_CLARO, 'B': AMARELO_CLARO, 'C': 'F2F2F2'}
    for p in produtos:
        ws.cell(row=linha, column=1, value=p['classe']).font = Font(bold=True, color=VERDE_ESCURO)
        ws.cell(row=linha, column=2, value=p['nome'])
        ws.cell(row=linha, column=3, value=p['categoria'])
        ws.cell(row=linha, column=4, value=round(p['qtd'], 3))
        ws.cell(row=linha, column=5, value=round(p['receita'], 2))
        ws.cell(row=linha, column=6, value=p['pct_receita'])
        ws.cell(row=linha, column=7, value=p['pct_acumulado'])
        for c in range(1, 8):
            ws.cell(row=linha, column=c).fill = PatternFill('solid', fgColor=cores[p['classe']])
        linha += 1
    ultima = linha - 1
    if produtos:
        _formatar_coluna(ws, 4, primeira, ultima, NUMERO)
        _formatar_coluna(ws, 5, primeira, ultima, MOEDA)
        _formatar_coluna(ws, 6, primeira, ultima, PCT)
        _formatar_coluna(ws, 7, primeira, ultima, PCT)
        ws.auto_filter.ref = f'A{primeira - 1}:G{ultima}'
        ws.freeze_panes = ws.cell(row=primeira, column=1)
        resumo = {}
        for p in produtos:
            resumo[p['classe']] = resumo.get(p['classe'], 0) + 1
        linha = ultima + 2
        ws.cell(row=linha, column=1, value='Resumo:').font = Font(bold=True)
        ws.cell(row=linha, column=2,
                value=' · '.join(f'Classe {k}: {v} produto(s)' for k, v in sorted(resumo.items())))
    else:
        ws.cell(row=primeira, column=2, value='Nenhuma venda no período.')
    return ws


def _aba_categorias(wb, inicio, fim, rotulo):
    dados = an.vendas_por_categoria(inicio, fim)
    ws = wb.create_sheet('Categorias')
    linha = _titulo_aba(ws, 'Vendas por categoria', f'Período: {rotulo}')
    linha = _cabecalho(ws, linha, ['Categoria', 'Quantidade', 'Receita', '% da receita', 'Nº de vendas'],
                       [26, 14, 16, 14, 14])
    primeira = linha
    for c in dados:
        ws.cell(row=linha, column=1, value=c['categoria'])
        ws.cell(row=linha, column=2, value=round(c['qtd'], 3))
        ws.cell(row=linha, column=3, value=c['receita'])
        ws.cell(row=linha, column=4, value=c['pct'])
        ws.cell(row=linha, column=5, value=c['n_vendas'])
        linha += 1
    ultima = linha - 1
    if dados:
        _formatar_coluna(ws, 2, primeira, ultima, NUMERO)
        _formatar_coluna(ws, 3, primeira, ultima, MOEDA)
        _formatar_coluna(ws, 4, primeira, ultima, PCT)
        _zebra(ws, primeira, ultima, 5)
        _grafico(ws, 'barra', 'Receita por categoria', 3, 1, primeira, ultima, f'G{primeira}')
    else:
        ws.cell(row=primeira, column=1, value='Nenhuma venda no período.')
    return ws


def _aba_pagamento(wb, inicio, fim, rotulo):
    dados = an.vendas_por_forma_pagamento(inicio, fim)
    ws = wb.create_sheet('Formas de Pagamento')
    linha = _titulo_aba(ws, 'Formas de pagamento', f'Período: {rotulo}')
    linha = _cabecalho(ws, linha, ['Forma', 'Vendas', 'Total', '% do total', 'Taxas pagas', 'Líquido recebido'],
                       [20, 12, 16, 13, 16, 18])
    primeira = linha
    for f in dados:
        ws.cell(row=linha, column=1, value=f.get('rotulo') or f['forma'])
        ws.cell(row=linha, column=2, value=f['n'])
        ws.cell(row=linha, column=3, value=f['total'])
        ws.cell(row=linha, column=4, value=f['pct'])
        ws.cell(row=linha, column=5, value=f['taxas'])
        ws.cell(row=linha, column=6, value=f['liquido'])
        linha += 1
    ultima = linha - 1
    if dados:
        for col in (3, 5, 6):
            _formatar_coluna(ws, col, primeira, ultima, MOEDA)
        _formatar_coluna(ws, 4, primeira, ultima, PCT)
        _zebra(ws, primeira, ultima, 6)
        _grafico(ws, 'barra', 'Total por forma de pagamento', 3, 1, primeira, ultima, f'H{primeira}')
    else:
        ws.cell(row=primeira, column=1, value='Nenhuma venda no período.')
    return ws


def _aba_tipo_dia(wb, inicio, fim, rotulo):
    classes = an.vendas_por_classe_dia(inicio, fim)
    ws = wb.create_sheet('Tipo de Dia')
    linha = _titulo_aba(
        ws, 'Feriado, véspera, dia de pagamento ou dia normal?',
        'O que de fato muda o movimento do mercadinho — não a data comemorativa em si. '
        f'Período: {rotulo}')
    linha = _cabecalho(ws, linha,
                       ['Tipo de dia', 'Dias no período', 'Vendas', 'Faturamento total',
                        'Faturamento médio por dia', 'Ticket médio'],
                       [22, 15, 12, 18, 22, 14])
    primeira = linha
    for c in classes:
        ws.cell(row=linha, column=1, value=c['classe'])
        ws.cell(row=linha, column=2, value=c['dias'])
        ws.cell(row=linha, column=3, value=c['numero_vendas'])
        ws.cell(row=linha, column=4, value=c['faturamento'])
        ws.cell(row=linha, column=5, value=c['media_por_dia'])
        ws.cell(row=linha, column=6, value=c['ticket_medio'])
        if c['classe'] in ('Feriado', 'Dia de pagamento') and c['dias']:
            for col in range(1, 7):
                ws.cell(row=linha, column=col).fill = PatternFill('solid', fgColor=AMARELO_CLARO)
        linha += 1
    ultima = linha - 1
    for col in (4, 5, 6):
        _formatar_coluna(ws, col, primeira, ultima, MOEDA)
    _grafico(ws, 'barra', 'Faturamento médio por dia, por tipo de dia', 5, 1, primeira, ultima, f'H{primeira}')
    linha = ultima + 16

    # Próximos feriados e dias de pagamento, pra planejar a compra da semana
    proximos = an.proximos_eventos(45)
    ws.cell(row=linha, column=1, value='PRÓXIMOS FERIADOS E DIAS DE PAGAMENTO').font = Font(bold=True, color=VERDE_ESCURO)
    linha += 1
    linha = _cabecalho(ws, linha, ['Data', 'Dia da semana', 'Tipo', 'Nome', 'Faltam (dias)'],
                       [12, 14, 18, 30, 14])
    primeira_prox = linha
    for e in proximos:
        ws.cell(row=linha, column=1, value=e['data_br'])
        ws.cell(row=linha, column=2, value=e['dia_semana'])
        ws.cell(row=linha, column=3, value=e['tipo'])
        ws.cell(row=linha, column=4, value=e['nome'])
        ws.cell(row=linha, column=5, value=e['dias_restantes'])
        linha += 1
    if proximos:
        _zebra(ws, primeira_prox, linha - 1, 5)
    else:
        ws.cell(row=linha, column=1, value='Nenhum feriado ou dia de pagamento nos próximos 45 dias.')
    return ws


def _aba_clientes(wb, inicio, fim, rotulo):
    clientes = an.top_clientes(inicio, fim, limite=100)
    ws = wb.create_sheet('Clientes')
    linha = _titulo_aba(ws, 'Clientes que mais compraram', f'Período: {rotulo}')
    if not clientes:
        ws.cell(row=linha, column=1,
                value='Nenhuma venda vinculada a cliente cadastrado neste período '
                      '(ou este banco não guarda o cliente da venda).')
        ws.column_dimensions['A'].width = 80
        return ws
    linha = _cabecalho(ws, linha, ['#', 'Cliente', 'Compras', 'Total gasto', 'Ticket médio', 'Última compra'],
                       [6, 34, 12, 16, 15, 16])
    primeira = linha
    for i, c in enumerate(clientes, start=1):
        ws.cell(row=linha, column=1, value=i)
        ws.cell(row=linha, column=2, value=c['nome'])
        ws.cell(row=linha, column=3, value=c['n_compras'])
        ws.cell(row=linha, column=4, value=c['total'])
        ws.cell(row=linha, column=5, value=c['ticket_medio'])
        ws.cell(row=linha, column=6, value=c['ultima_compra'])
        linha += 1
    ultima = linha - 1
    for col in (4, 5):
        _formatar_coluna(ws, col, primeira, ultima, MOEDA)
    _zebra(ws, primeira, ultima, 6)
    ws.freeze_panes = ws.cell(row=primeira, column=1)
    return ws


# ----------------------------------------------------------------------
# Função principal
# ----------------------------------------------------------------------
def gerar_relatorio_xlsx(caminho_destino, inicio=None, fim=None):
    """Gera a planilha completa de análises no caminho informado.
    inicio/fim no formato 'AAAA-MM-DD'; None em ambos = todo o histórico."""
    if openpyxl is None:
        raise RuntimeError('Este recurso precisa do pacote openpyxl. Rode: pip install openpyxl')

    rotulo = an.rotulo_periodo(inicio, fim)
    wb = openpyxl.Workbook()

    _aba_resumo(wb, inicio, fim, rotulo)
    _aba_serie(wb, 'Evolução por Dia', 'dia', 'Evolução por dia', inicio, fim, rotulo, 'linha')
    _aba_serie(wb, 'Evolução por Semana', 'semana', 'Evolução por semana', inicio, fim, rotulo, 'barra')
    _aba_serie(wb, 'Evolução por Mês', 'mes', 'Evolução por mês', inicio, fim, rotulo, 'barra')
    _aba_dia_semana(wb, inicio, fim, rotulo)
    _aba_produtos_por_dia(wb, inicio, fim, rotulo)
    _aba_horario(wb, inicio, fim, rotulo)
    _aba_top_produtos(wb, inicio, fim, rotulo)
    _aba_curva_abc(wb, inicio, fim, rotulo)
    _aba_categorias(wb, inicio, fim, rotulo)
    _aba_pagamento(wb, inicio, fim, rotulo)
    _aba_tipo_dia(wb, inicio, fim, rotulo)
    _aba_clientes(wb, inicio, fim, rotulo)

    wb.save(caminho_destino)
    resumo = an.resumo_periodo(inicio, fim)
    return {
        'caminho': caminho_destino,
        'periodo': rotulo,
        'abas': len(wb.sheetnames),
        'faturamento': resumo['faturamento'],
        'numero_vendas': resumo['numero_vendas'],
    }


if __name__ == '__main__':
    destino = sys.argv[1] if len(sys.argv) > 1 else 'relatorio_vendas.xlsx'
    ini = sys.argv[2] if len(sys.argv) > 2 else None
    fim_ = sys.argv[3] if len(sys.argv) > 3 else None
    db.init_db()
    r = gerar_relatorio_xlsx(destino, ini, fim_)
    print(f"Relatório gerado: {r['caminho']}")
    print(f"Período: {r['periodo']} | {r['numero_vendas']} venda(s) | R$ {r['faturamento']:.2f}")
    print(f"Abas criadas: {r['abas']}")
