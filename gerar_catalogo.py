"""
gerar_catalogo.py

Gera (ou regenera) o "Catálogo de Produtos" em .xlsx direto a partir do
mercadinho.db em uso — a mesma ideia da planilha que já existia pra bipar
os itens de balança na tela (código entre asteriscos com a fonte
"Libre Barcode 39"), só que agora sempre sincronizada com o que está
realmente cadastrado no sistema, em vez de mantida à mão.

Uso pela interface: Arquivo -> Gerar catálogo atualizado (.xlsx).
Uso direto (linha de comando): python gerar_catalogo.py [caminho_saida.xlsx]

Se a fonte "Libre Barcode 39" (gratuita, disponível no Google Fonts) não
estiver instalada no Windows que for abrir a planilha, a coluna "Cód.
Barras (tela)" continua funcionando como texto normal — só não vira
desenho de código de barras na tela até a fonte ser instalada.
"""
from __future__ import annotations

import sys
from datetime import datetime

import database as db

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    openpyxl = None

CORES_CATEGORIA = ['2E7D32', '1565C0', 'AD1457', 'E65100', '4527A0', '00838F', '5D4037', '37474F']
COR_LINHA_BALANCA = 'FFF9C4'
COR_CABECALHO = 'D9D9D9'
FONTE_BARRAS = 'Libre Barcode 39'
CATEGORIA_SEM = 'Sem categoria'


def _proximo_codigo_plu(usados, inicio=9000):
    n = inicio
    while str(n) in usados:
        n += 1
    usados.add(str(n))
    return str(n)


def gerar_catalogo_xlsx(caminho_destino, preencher_codigos_faltantes=True):
    """Lê os produtos ativos do banco, opcionalmente preenche um código novo
    pros itens de balança que ainda não têm nenhum (só preenche quem está
    vazio — nunca troca um código que já existe), e escreve a planilha.

    Devolve um dict com o resumo do que foi feito, pra interface mostrar."""
    import gestao as _gestao
    _gestao.permitir('estoque')
    if openpyxl is None:
        raise RuntimeError('Este recurso precisa do pacote openpyxl. Rode: pip install openpyxl')

    conn = db.get_conn()
    try:
        produtos = [dict(r) for r in conn.execute(
            """SELECT id, codigo_barras, nome, categoria, unidade, preco_venda
               FROM produtos WHERE ativo = 1 ORDER BY COALESCE(categoria, 'zzz'), nome""")]

        codigos_preenchidos = []
        if preencher_codigos_faltantes:
            usados = {p['codigo_barras'] for p in produtos if p['codigo_barras']}
            for p in produtos:
                sem_codigo = not (p['codigo_barras'] and str(p['codigo_barras']).strip())
                if p['categoria'] == 'Balança' and sem_codigo:
                    novo = _proximo_codigo_plu(usados)
                    conn.execute('UPDATE produtos SET codigo_barras = ? WHERE id = ?', (novo, p['id']))
                    p['codigo_barras'] = novo
                    codigos_preenchidos.append({'nome': p['nome'], 'codigo': novo})
            if codigos_preenchidos:
                conn.commit()

        sem_categoria = [p for p in produtos if not (p['categoria'] and p['categoria'].strip())]
        for p in produtos:
            if not (p['categoria'] and p['categoria'].strip()):
                p['categoria'] = CATEGORIA_SEM
    finally:
        conn.close()

    wb = openpyxl.Workbook()

    # --- Índice ---
    idx = wb.active
    idx.title = 'Índice'
    idx['A1'] = 'Catálogo de Produtos — FMG Mercadinho'
    idx['A1'].font = Font(size=16, bold=True, color='1B5240')
    idx['A3'] = 'Gerado automaticamente a partir do mercadinho.db em uso — não precisa mais editar essa planilha à mão.'
    idx['A5'] = 'Total de produtos ativos'
    idx['B5'] = len(produtos)
    idx['A6'] = 'Gerado em'
    idx['B6'] = datetime.now().strftime('%d/%m/%Y %H:%M')
    linha = 8
    if codigos_preenchidos:
        idx[f'A{linha}'] = f'{len(codigos_preenchidos)} produto(s) de balança ganharam um código novo (não tinham nenhum):'
        idx[f'A{linha}'].font = Font(bold=True, color='9A6700')
        linha += 1
        for c in codigos_preenchidos:
            idx[f'A{linha}'] = c['nome']
            idx[f'B{linha}'] = c['codigo']
            linha += 1
        linha += 1
    if sem_categoria:
        idx[f'A{linha}'] = f'{len(sem_categoria)} produto(s) sem categoria definida (revise na aba Produtos):'
        idx[f'A{linha}'].font = Font(bold=True, color='C23B2E')
        linha += 1
        for p in sem_categoria[:40]:
            idx[f'A{linha}'] = p['nome']
            linha += 1
    idx.column_dimensions['A'].width = 55
    idx.column_dimensions['B'].width = 20

    # --- Catálogo (agrupado por categoria, pronto pra imprimir/bipar da tela) ---
    cat = wb.create_sheet('Catálogo')
    cat['A1'] = ('Itens de BALANÇA (destacados em amarelo) — aponte o leitor pra célula da coluna '
                 '"Cód. Barras (tela)" e bipe direto na tela.')
    for col, largura in zip('ABCDE', [16, 44, 10, 14, 26]):
        cat.column_dimensions[col].width = largura
    linha_atual = 3
    categorias_ordenadas = sorted({p['categoria'] for p in produtos})
    for i, categoria in enumerate(categorias_ordenadas):
        itens = [p for p in produtos if p['categoria'] == categoria]
        cor = CORES_CATEGORIA[i % len(CORES_CATEGORIA)]
        cat.cell(row=linha_atual, column=1, value=f'{categoria}  ({len(itens)} itens)').font = Font(bold=True, size=12, color=cor)
        linha_atual += 1
        for c, titulo in enumerate(['Código', 'Produto', 'Unidade', 'Valor (R$)', 'Cód. Barras (tela)'], start=1):
            cel = cat.cell(row=linha_atual, column=c, value=titulo)
            cel.font = Font(bold=True)
            cel.fill = PatternFill('solid', fgColor=COR_CABECALHO)
        linha_atual += 1
        for p in itens:
            eh_balanca = p['categoria'] == 'Balança'
            cat.cell(row=linha_atual, column=1, value=p['codigo_barras'] or '')
            cat.cell(row=linha_atual, column=2, value=p['nome'])
            cat.cell(row=linha_atual, column=3, value=p['unidade'])
            cat.cell(row=linha_atual, column=4, value=p['preco_venda'])
            f_cel = cat.cell(row=linha_atual, column=5, value=f'=IF(A{linha_atual}="","","*"&A{linha_atual}&"*")')
            f_cel.font = Font(name=FONTE_BARRAS, size=13)
            if eh_balanca:
                for c in range(1, 6):
                    cat.cell(row=linha_atual, column=c).fill = PatternFill('solid', fgColor=COR_LINHA_BALANCA)
            linha_atual += 1
        linha_atual += 1

    # --- Lista Completa (Filtrável) ---
    lista = wb.create_sheet('Lista Completa (Filtrável)')
    for c, titulo in enumerate(['Categoria', 'Código', 'Produto', 'Unidade', 'Valor (R$)', 'Cód. Barras (tela)'], start=1):
        cel = lista.cell(row=1, column=c, value=titulo)
        cel.font = Font(bold=True)
        cel.fill = PatternFill('solid', fgColor=COR_CABECALHO)
    for r, p in enumerate(produtos, start=2):
        lista.cell(row=r, column=1, value=p['categoria'])
        lista.cell(row=r, column=2, value=p['codigo_barras'] or '')
        lista.cell(row=r, column=3, value=p['nome'])
        lista.cell(row=r, column=4, value=p['unidade'])
        lista.cell(row=r, column=5, value=p['preco_venda'])
        f_cel = lista.cell(row=r, column=6, value=f'=IF(B{r}="","","*"&B{r}&"*")')
        f_cel.font = Font(name=FONTE_BARRAS, size=13)
    lista.auto_filter.ref = f'A1:F{len(produtos) + 1}'
    for col, largura in zip('ABCDEF', [22, 16, 44, 10, 12, 24]):
        lista.column_dimensions[col].width = largura
    lista.freeze_panes = 'A2'

    # --- Como usar ---
    ajuda = wb.create_sheet('Como usar o código de barras')
    linhas_ajuda = [
        'COMO FUNCIONA A COLUNA "CÓD. BARRAS (TELA)"',
        '',
        'Ela pega o código do produto e coloca um asterisco (*) antes e depois '
        '(ex.: 9001 vira *9001*) — é o padrão do formato Code 39. Com a fonte '
        '"Libre Barcode 39" aplicada (gratuita, Google Fonts), esse texto vira '
        'um código de barras de verdade na tela.',
        '',
        'Aponte o leitor USB pra essa célula e bipe — funciona igual a um código impresso.',
        '',
        'Isso ajuda sobretudo os itens de BALANÇA (frutas, verduras, pães soltos) '
        'que não vêm com um código de barras de fábrica.',
        '',
        'Sempre que cadastrar produto novo ou mudar preço, é só gerar essa planilha '
        'de novo pelo menu Arquivo -> Gerar catálogo atualizado.',
    ]
    for i, texto in enumerate(linhas_ajuda, start=1):
        ajuda.cell(row=i, column=1, value=texto)
    ajuda.column_dimensions['A'].width = 100

    wb.save(caminho_destino)
    return {
        'total_produtos': len(produtos),
        'codigos_preenchidos': codigos_preenchidos,
        'sem_categoria': len(sem_categoria),
        'caminho': caminho_destino,
    }


if __name__ == '__main__':
    destino = sys.argv[1] if len(sys.argv) > 1 else 'catalogo_gerado.xlsx'
    db.init_db()
    resultado = gerar_catalogo_xlsx(destino)
    print(f"Catálogo gerado em: {resultado['caminho']}")
    print(f"Produtos exportados: {resultado['total_produtos']}")
    if resultado['codigos_preenchidos']:
        print(f"Códigos novos criados para itens de balança sem código: {len(resultado['codigos_preenchidos'])}")
    if resultado['sem_categoria']:
        print(f"Atenção: {resultado['sem_categoria']} produto(s) sem categoria definida.")
